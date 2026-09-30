from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

source = Path(sys.argv[1])
out = Path(sys.argv[2])
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.8)
section.bottom_margin = Inches(.8)
section.left_margin = Inches(.9)
section.right_margin = Inches(.9)
section.header_distance = Inches(.35)
section.footer_distance = Inches(.35)

for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.underline = False
    for el in list(style.element.xpath('.//w:pBdr')):
        el.getparent().remove(el)

normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True

title = doc.styles['Title']
title.font.size = Pt(25)
title.font.bold = True
title.paragraph_format.space_after = Pt(15)
title.paragraph_format.line_spacing = 1.05

for name, size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
    style=doc.styles[name]
    style.font.size=Pt(size)
    style.font.bold=True
    style.paragraph_format.space_before=Pt(13)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.keep_with_next=True

code = doc.styles.add_style('Article Code', 1)
code.font.name = 'DejaVu Sans Mono'
code.font.size = Pt(8.5)
code.font.color.rgb = RGBColor(0,0,0)
code.paragraph_format.line_spacing = 1.05
code.paragraph_format.space_before = Pt(3)
code.paragraph_format.space_after = Pt(10)
code.paragraph_format.keep_together = True
code.paragraph_format.left_indent = Inches(.08)

# Keep this as an editable article rather than a cover-page report.
footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run()
run.font.name = 'Arial'
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(90,90,90)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
run._r.addnext(fld)


def inline(paragraph, text):
    for i, part in enumerate(re.split(r'(`[^`]+`)', text)):
        r=paragraph.add_run(part[1:-1] if part.startswith('`') and part.endswith('`') else part)
        if part.startswith('`') and part.endswith('`'):
            r.font.name='DejaVu Sans Mono'
            r.font.size=Pt(9.5)

blocks=source.read_text().split('\n\n')
for block in blocks:
    block=block.strip()
    if not block:
        continue
    if block.startswith('```'):
        lines=block.splitlines()[1:]
        if lines and lines[-1]=='```': lines.pop()
        doc.add_paragraph('\n'.join(lines), 'Article Code')
    elif block.startswith('# '):
        doc.add_paragraph(block[2:], 'Title')
    elif block.startswith('## '):
        doc.add_paragraph(block[3:], 'Heading 1')
    else:
        p=doc.add_paragraph()
        inline(p, block.replace('\n',' '))

doc.core_properties.title='How I connected dot to iMessage with Photon'
doc.core_properties.subject='Editable long-form article draft'
doc.core_properties.author=''
doc.core_properties.keywords='Photon, iMessage, dot, architecture'
doc.save(out)
print(out)
